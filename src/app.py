import streamlit as st
import pandas as pd
import docx
import io
import openpyxl
from typing import List, Dict, Optional
from pathlib import Path
from converter import MessageConverter, ProcessingError
from config import Config
import time

class DocxConverterApp:
    def __init__(self):
        self.converter = MessageConverter()
        self.config = Config()
        self.setup_streamlit()

    def setup_streamlit(self):
        """Streamlit設定の初期化"""
        st.set_page_config(
            page_title="DOCX Converter",
            page_icon="📄",
            layout="wide"
        )
        # セッション状態の初期化
        if 'processed_df' not in st.session_state:
            st.session_state.processed_df = None
        if 'errors' not in st.session_state:
            st.session_state.errors = []
        if 'selected_columns' not in st.session_state:
            st.session_state.selected_columns = []
        if 'available_columns' not in st.session_state:
            st.session_state.available_columns = set()

    def read_docx(self, file) -> Dict[str, str]:
        """DOCXファイルの読み込み"""
        try:
            doc = docx.Document(file)
            content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return {"document_content": content}
        except Exception as e:
            raise ValueError(f"ファイルの読み込みに失敗しました: {str(e)}")

    def create_excel_with_highlights(self, df: pd.DataFrame, errors: List[ProcessingError]) -> bytes:
        """エラーハイライト付きExcelの作成"""
        try:
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='データ')
                workbook = writer.book
                worksheet = writer.sheets['データ']
                
                # エラー行の背景色を設定
                error_rows = set(error.document_index for error in errors)
                for row_idx in error_rows:
                    for cell in worksheet[row_idx + 2]:  # Excel行は1から始まり、ヘッダーがあるため+2
                        cell.fill = openpyxl.styles.PatternFill(
                            start_color='FFE7E6',
                            end_color='FFE7E6',
                            fill_type='solid'
                        )
                
                # カラム幅の自動調整
                for column in worksheet.columns:
                    max_length = 0
                    column_name = column[0].value
                    
                    # すべての行をチェックして最大長を取得
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    
                    # 最大長に基づいてカラム幅を設定（最小幅10、最大幅50）
                    adjusted_width = min(max(10, max_length + 2), 50)
                    worksheet.column_dimensions[column[0].column_letter].width = adjusted_width
            
            return output.getvalue()
        except Exception as e:
            raise ValueError(f"Excelファイルの作成に失敗しました: {str(e)}")

    def display_preview(self, content: str):
        """文書のプレビュー表示"""
        st.subheader("プレビュー")
        with st.expander("文書内容", expanded=True):
            st.text_area("", value=content[:500] + "..." if len(content) > 500 else content, 
                        height=200, disabled=True)

    def display_errors(self, errors: List[ProcessingError]):
        """エラーの表示"""
        if errors:
            st.error("処理中にエラーが発生しました")
            for error in errors:
                with st.expander(f"文書 {error.document_index + 1} のエラー"):
                    st.write(f"種類: {error.error_type}")
                    st.write(f"メッセージ: {error.message}")
                    if error.details:
                        st.write(f"詳細: {error.details}")

    def get_default_columns(self, available_columns: set) -> List[str]:
        """デフォルトのカラムを取得"""
        default_columns = []
        if available_columns:
            default_columns = ['番号']  # 番号は常に含める
            # 優先順位に基づいてカラムを追加
            priority_columns = ['原稿', '部署', '名前']
            for col in priority_columns:
                if col in available_columns:
                    default_columns.append(col)
        return default_columns

    def process_files(self, files: List) -> Optional[List[Dict]]:
        """ファイルの処理"""
        documents = []
        progress_bar = st.progress(0)
        status_text = st.empty()

        try:
            for i, file in enumerate(files):
                # 進捗表示の更新
                progress = (i + 1) / len(files)
                progress_bar.progress(progress)
                status_text.text(f"処理中... {i + 1}/{len(files)} ファイル")

                doc = self.read_docx(file)
                documents.append(doc)
                time.sleep(0.1)  # UI更新のための小さな遅延

            status_text.text("処理完了")
            return documents

        except Exception as e:
            progress_bar.empty()
            status_text.empty()
            st.error(f"ファイル処理中にエラーが発生しました: {str(e)}")
            return None

    def run(self):
        """アプリケーションのメイン処理"""
        st.title("文書変換アプリ")

        uploaded_files = st.file_uploader(
            "DOCXファイルをアップロードしてください",
            type="docx",
            accept_multiple_files=True
        )

        if uploaded_files:
            documents = self.process_files(uploaded_files)
            
            if documents:
                try:
                    # 利用可能なカラムを取得
                    st.session_state.available_columns = \
                        self.converter.extract_potential_columns(documents[0]['document_content'])
                    
                    # プレビュー表示
                    self.display_preview(documents[0]['document_content'])
                    
                    # カラム選択UI
                    col1, col2 = st.columns(2)
                    with col1:
                        # デフォルトのカラムを取得
                        default_columns = self.get_default_columns(st.session_state.available_columns)
                        
                        st.session_state.selected_columns = st.multiselect(
                            "使用するカラムを選択：",
                            list(st.session_state.available_columns),
                            default=default_columns
                        )
                    
                    with col2:
                        if st.button("カラム設定を保存"):
                            self.config.update_favorite_columns(st.session_state.selected_columns)
                            st.success("カラム設定を保存しました")

                    if st.session_state.selected_columns:
                        self.converter.selected_columns = st.session_state.selected_columns

                        # データ処理
                        df, errors = self.converter.process_documents(documents)
                        
                        # エラー表示
                        self.display_errors(errors)

                        # 結果表示
                        if not df.empty:
                            st.subheader("変換結果")
                            
                            # データフレームの編集機能
                            edited_df = st.data_editor(
                                df,
                                num_rows="dynamic",
                                use_container_width=True,
                                height=400
                            )
                            
                            # エクセルファイルのダウンロード
                            excel_data = self.create_excel_with_highlights(edited_df, errors)
                            st.download_button(
                                label="Excelファイルをダウンロード",
                                data=excel_data,
                                file_name="converted_data.xlsx",
                                mime="application/vnd.ms-excel",
                                key='download_button'
                            )

                except Exception as e:
                    st.error(f"予期せぬエラーが発生しました: {str(e)}")

if __name__ == "__main__":
    app = DocxConverterApp()
    app.run()
